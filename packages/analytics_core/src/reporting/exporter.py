"""Structured Document Export Engine for MindEd AA-OS (PDF and DOCX).

Renders the canonical structured analytical answer with:
- MindEd Emblem / Logo and Cryptographic Provenance Fingerprint
- Canonical SHA-256 Content Digest & Verifiable Document Seal
- Status, Claim & Epistemic Verification Badges
- Headline (One sentence answer)
- Key Result (Primary metric value, sample size, observation counts)
- Task-Specific Structured Table (Rankings, Statistical associations, Group comparisons)
- Key Insights (crisp one-line insights)
- Limitations & Assumptions (epistemic bounds, non-causal caveats)
- Recommended Actions (numbered operational recommendations)
- Audited Evidence & Provenance (calculations, rows analyzed, SQL queries, SHA-256 manifest)
"""

from io import BytesIO
import hashlib
import html
import json
import re
from datetime import datetime, timezone

from typing import Any, Dict, List, Optional, Tuple

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from PIL import Image, ImageDraw

from packages.analytics_core.src.reporting.narrative_composer import compose_report_narrative

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
    Image as RLImage,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch


def _fmt_compact(val: Optional[float]) -> str:
    if val is None:
        return "N/A"
    try:
        val = float(val)
        if abs(val) >= 1_000_000:
            return f"{val / 1_000_000:.2f}M"
        if abs(val) >= 1_000:
            return f"{val:,.2f}"
        return f"{val:,.4f}"
    except (ValueError, TypeError):
        return str(val)


def _as_dict(value: Any) -> Dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _finite_number(value: Any) -> bool:
    try:
        number = float(value)
        return number == number and number not in (float("inf"), float("-inf"))
    except (TypeError, ValueError, OverflowError):
        return False


def _json_safe(value: Any) -> Any:
    """Convert report metadata into deterministic JSON-compatible values."""
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if _finite_number(value) else str(value)
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_safe(v) for v in value]
    if hasattr(value, "item"):
        try:
            return _json_safe(value.item())
        except Exception:
            pass
    return str(value)


def _normalise_chart_candidate(value: Any, index: int = 0) -> Optional[Dict[str, Any]]:
    """Validate a saved chart payload without generating or filling datapoints."""
    if not isinstance(value, dict):
        return None
    raw_type = str(value.get("chart_type") or value.get("type") or "").strip().lower()
    aliases = {"pie": "donut", "timeseries": "line", "time_series": "line", "line_chart": "line"}
    chart_type = aliases.get(raw_type, raw_type)
    if chart_type not in {"line", "multi_line", "bar", "grouped_bar", "stacked_bar", "scatter", "histogram", "donut", "boxplot", "waterfall"}:
        return None
    raw_data = value.get("data")
    if not isinstance(raw_data, list) or not raw_data or not all(isinstance(row, dict) for row in raw_data):
        return None
    data = [_json_safe(row) for row in raw_data]
    x_axis = _as_dict(value.get("x_axis"))
    y_axis = _as_dict(value.get("y_axis"))
    first = data[0]
    x_field = str(x_axis.get("field") or value.get("x_field") or value.get("xField") or next(iter(first), "category"))
    y_field = str(y_axis.get("field") or value.get("y_field") or value.get("yField") or next((k for k, v in first.items() if _finite_number(v)), "value"))
    if chart_type not in {"boxplot", "waterfall"} and not all(
        row.get(x_field) is not None and _finite_number(row.get(y_field)) for row in data
    ):
        return None
    if chart_type == "boxplot":
        required = ("p10", "p25", "p50", "p75", "p90")
        for row in data:
            if not all(_finite_number(row.get(field)) for field in required):
                return None
            if not (float(row["p10"]) <= float(row["p25"]) <= float(row["p50"]) <= float(row["p75"]) <= float(row["p90"])):
                return None
    if chart_type == "waterfall" and not all(
        _finite_number(row.get("delta", row.get(y_field)))
        and _finite_number(row.get("start", 0))
        and _finite_number(row.get("end", row.get("subtotal", row.get("delta", row.get(y_field)))))
        for row in data
    ):
        return None
    return {
        "chart_id": str(value.get("chart_id") or f"chart_{index + 1}"),
        "chart_type": chart_type,
        "title": str(value.get("title") or value.get("name") or f"{y_field} by {x_field}"),
        "explanation": str(value.get("explanation") or ""),
        "x_axis": {"field": x_field, "label": str(x_axis.get("label") or x_axis.get("title") or x_field), "type": str(x_axis.get("type") or "")},
        "y_axis": {"field": y_field, "label": str(y_axis.get("label") or y_axis.get("title") or y_field), "type": str(y_axis.get("type") or ""), "format": str(y_axis.get("format") or "")},
        "series_field": str(value.get("series_field") or value.get("seriesField") or ""),
        "data": data,
        "metadata": _json_safe(_as_dict(value.get("metadata"))),
    }


def _select_report_visualizations(investigation_data: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Create chart specs only from saved output data; never re-run analysis or invent rows."""
    result = _as_dict(investigation_data.get("analyst_result"))
    numbers = _as_dict(result.get("numbers"))
    question = str(investigation_data.get("question") or "")
    kind = str(result.get("kind") or "ANALYSIS").upper()
    status = str(investigation_data.get("status") or "").upper()
    evidence = investigation_data.get("evidence") or []
    execution = _as_dict(investigation_data.get("execution_record") or numbers.get("execution_record"))
    if status != "COMPLETED" or not evidence:
        return []
    # Match the browser's verification gate: no confident chart if any required
    # evidence row failed or was never independently verified.
    if not all(str(_as_dict(ev).get("validation_status") or "").upper() in {"VERIFIED", "PASSED"} for ev in evidence):
        return []
    if execution and (str(execution.get("contract_status") or "").upper() != "MATCHED" or str(execution.get("verification_status") or "").upper() != "VERIFIED"):
        return []
    if not execution and not evidence:
        return []

    explicit_candidates: List[Any] = []
    for key in ("trend_charts", "chart_specs", "visualizations"):
        if isinstance(numbers.get(key), list):
            explicit_candidates.extend(numbers[key])
    explicit_candidates.extend([
        numbers.get("trend_chart"), numbers.get("trend_slope_chart"),
        numbers.get("forecast_chart"), numbers.get("comparison_chart"),
        numbers.get("distribution_chart"), numbers.get("relationship_chart"),
        numbers.get("scatter_chart"),
    ])
    explicit = []
    seen = set()
    for index, candidate in enumerate(explicit_candidates):
        chart = _normalise_chart_candidate(candidate, index)
        if not chart:
            continue
        identity = (chart["chart_id"], chart["chart_type"], chart["title"])
        if identity not in seen:
            explicit.append(chart)
            seen.add(identity)
    if explicit:
        return explicit

    operation = str(numbers.get("operation") or numbers.get("aggregation") or _as_dict(numbers.get("canonical_contract")).get("requested_aggregation") or "").upper()
    metric = str(numbers.get("target_column") or _as_dict(numbers.get("canonical_contract")).get("target_column") or "metric")
    grouping = numbers.get("grouping_columns") or _as_dict(numbers.get("canonical_contract")).get("grouping_columns") or []
    group_field = str(grouping[0]) if isinstance(grouping, list) and grouping else "category"
    charts: List[Dict[str, Any]] = []

    def make_chart(chart_id: str, chart_type: str, title: str, x_field: str, y_field: str, rows: List[Dict[str, Any]], explanation: str = "", metadata: Optional[Dict[str, Any]] = None, series_field: str = "") -> Dict[str, Any]:
        return {
            "chart_id": chart_id, "chart_type": chart_type, "title": title, "explanation": explanation,
            "x_axis": {"field": x_field, "label": x_field, "type": "nominal" if chart_type in {"bar", "grouped_bar", "stacked_bar", "histogram", "donut", "boxplot"} else "quantitative"},
            "y_axis": {"field": y_field, "label": y_field, "type": "quantitative"},
            "series_field": series_field, "data": _json_safe(rows), "metadata": _json_safe(metadata or {"source_method": "persisted_analyst_result"}),
        }

    composition_spec = _as_dict(numbers.get("composition_over_time"))
    composition_data = composition_spec.get("data")
    time_field = str(composition_spec.get("time_field") or "")
    series_field = str(composition_spec.get("series_field") or "")
    value_field = str(composition_spec.get("value_field") or "")
    if (isinstance(composition_data, list) and len(composition_data) >= 2 and time_field and series_field and value_field
            and all(isinstance(row, dict) and row.get(time_field) is not None and isinstance(row.get(series_field), str) and _finite_number(row.get(value_field)) for row in composition_data)):
        charts.append(make_chart(
            f"composition_over_time_{value_field}", "stacked_bar",
            f"Composition of {value_field} over time by {series_field}",
            time_field, value_field, composition_data,
            "Stacked segments use saved period-by-group aggregates. The visualization does not invent missing periods or interpolate values.",
            {"source_method": "persisted_composition_over_time", "all_saved_points_included": True},
            series_field=series_field,
        ))
        return charts

    # Relationship charts require persisted paired raw observations. A correlation
    # coefficient or two marginal summaries are not enough to reconstruct a scatterplot.
    scatter_payload = _as_dict(numbers.get("scatter_data") or numbers.get("paired_observations"))
    scatter_raw = numbers.get("scatter_data") if isinstance(numbers.get("scatter_data"), list) else numbers.get("paired_observations") if isinstance(numbers.get("paired_observations"), list) else scatter_payload.get("data", scatter_payload.get("points", []))
    scatter_x = str(scatter_payload.get("x_field") or scatter_payload.get("x_column") or numbers.get("scatter_x_field") or numbers.get("explanatory_column") or numbers.get("secondary_target_column") or "")
    scatter_y = str(scatter_payload.get("y_field") or scatter_payload.get("y_column") or numbers.get("scatter_y_field") or numbers.get("target_column") or "")
    if isinstance(scatter_raw, list) and len(scatter_raw) >= 2 and scatter_x and scatter_y and all(
        isinstance(row, dict) and _finite_number(row.get(scatter_x)) and _finite_number(row.get(scatter_y)) for row in scatter_raw
    ):
        charts.append(make_chart(
            f"scatter_{scatter_x}_{scatter_y}", "scatter",
            f"Relationship between {scatter_x} and {scatter_y}", scatter_x, scatter_y,
            [{scatter_x: float(row[scatter_x]), scatter_y: float(row[scatter_y])} for row in scatter_raw],
            "Each point is a saved paired observation. The scatter plot shows association and does not establish causation.",
            {"source_method": "persisted_paired_observations", "all_saved_points_included": True},
        ))

    asks_distribution = bool(re.search(r"distribution|variability|spread|outlier|dispersion|volatility|consistent|range of values", question, re.I))
    histogram_payload = _as_dict(numbers.get("histogram"))
    histogram_raw = numbers.get("histogram_bins") if isinstance(numbers.get("histogram_bins"), list) else histogram_payload.get("bins", [])
    histogram_rows: List[Dict[str, Any]] = []
    if asks_distribution and isinstance(histogram_raw, list):
        for row in histogram_raw:
            row = _as_dict(row)
            start = row.get("bin_start", row.get("start", row.get("left")))
            end = row.get("bin_end", row.get("end", row.get("right")))
            count = row.get("count", row.get("frequency", row.get("n")))
            if _finite_number(start) and _finite_number(end) and _finite_number(count) and float(end) >= float(start) and float(count) >= 0:
                label = str(row.get("bin_label") or row.get("label") or f"{_fmt_compact(float(start))}–{_fmt_compact(float(end))}")
                histogram_rows.append({"bin_label": label, "count": float(count)})
    if histogram_rows:
        charts.append(make_chart(
            f"histogram_{metric}", "histogram", f"Distribution of {metric}", "bin_label", "count", histogram_rows,
            "Histogram bars use saved bin boundaries and frequencies; raw observations are not reconstructed.",
            {"source_method": "persisted_histogram_bins", "bin_count": len(histogram_rows)},
        ))

    table = _as_dict(numbers.get("table"))
    table_entries = []
    for label, raw in table.items():
        row = _as_dict(raw)
        value = row.get("value")
        if not _finite_number(value):
            table_entries = []
            break
        table_entries.append({"label": str(label), "value": float(value), "n": row.get("n")})
    if len(table_entries) >= 2:
        asc = bool(re.search(r"\b(lowest|least|smallest|ascending)\b", question, re.I) and not re.search(r"\b(highest|largest|top|descending)\b", question, re.I))
        ordered = sorted(table_entries, key=lambda row: row["value"], reverse=not asc)
        chart_entries = ordered[:50]
        bar_rows = [{group_field: row["label"], metric: row["value"], **({"n": row["n"]} if _finite_number(row.get("n")) else {})} for row in chart_entries]
        explanation = "Bars use saved group-level results; no separate model was fitted for this visualization."
        if len(ordered) > len(chart_entries):
            explanation += f" Showing {len(chart_entries)} groups of {len(ordered)} returned groups; the report table remains the full result summary."
        charts.append(make_chart(f"comparison_{group_field}_{metric}", "bar", f"{operation or 'Observed value'} of {metric} by {group_field}", group_field, metric, bar_rows,
                                 explanation, {"source_method": "persisted_aggregate_table", "ranking_order": "ascending" if asc else "descending", "groups_available": len(ordered), "groups_shown": len(chart_entries)}))
        total = sum(row["value"] for row in ordered)
        if operation in {"SUM", "COUNT", "ROW_COUNT"} and len(ordered) <= 8 and all(row["value"] >= 0 for row in ordered) and total > 0:
            charts.append(make_chart(f"composition_{group_field}_{metric}", "donut", f"Share of returned total {metric} by {group_field}", group_field, metric,
                                     [{group_field: row["label"], metric: row["value"]} for row in ordered],
                                     "Share of the returned non-negative additive group totals, not an independent re-aggregation.", {"source_method": "persisted_aggregate_table", "composition_total": total}))
        return charts

    for key, title, typ in (("means", "Mean by group", "bar"), ("rates", "Rate by group", "bar")):
        values = _as_dict(numbers.get(key))
        entries = [(str(k), float(v)) for k, v in values.items() if _finite_number(v)]
        if len(entries) >= 2:
            entries.sort(key=lambda item: item[1], reverse=True)
            is_rate = key == "rates"
            charts.append(make_chart(f"{key}_{metric}", typ, f"{title}: {metric}", "category", "value",
                                     [{"category": name, "value": value} for name, value in entries], metadata={"source_method": f"persisted_{key}"}))
            return charts

    correlations = _as_dict(numbers.get("correlations"))
    correlations_ok = [(str(k), float(v)) for k, v in correlations.items() if _finite_number(v) and abs(float(v)) <= 1]
    if len(correlations_ok) >= 2:
        correlations_ok.sort(key=lambda item: abs(item[1]), reverse=True)
        correlations_shown = correlations_ok[:50]
        corr_explanation = "Correlation describes linear association; it does not establish causation or per-predictor significance."
        if len(correlations_ok) > len(correlations_shown):
            corr_explanation += f" Showing the 50 largest absolute correlations of {len(correlations_ok)} recorded coefficients."
        charts.append(make_chart(f"correlations_{metric}", "bar", f"Recorded correlations with {metric}", "variable", "correlation",
                                 [{"variable": name, "correlation": value} for name, value in correlations_shown],
                                 corr_explanation, {"source_method": "persisted_correlation_coefficients", "coefficients_available": len(correlations_ok), "coefficients_shown": len(correlations_shown)}))
        return charts

    profile = _as_dict(numbers.get("distribution_profile"))
    if asks_distribution and profile:
        profiles = []
        for label, p in (("Prior period", _as_dict(profile.get("prior_period"))), ("Current period", _as_dict(profile.get("current_period")))):
            if p:
                profiles.append((label, p))
        if not profiles:
            profiles = [("Observed distribution", profile)]
        data = []
        whisker_definitions = set()
        for label, p in profiles:
            has_percentile_whiskers = _finite_number(p.get("p10")) and _finite_number(p.get("p90"))
            fields = {"p10": p.get("p10", p.get("min")), "p25": p.get("p25", p.get("q1")),
                      "p50": p.get("p50", p.get("median")), "p75": p.get("p75", p.get("q3")), "p90": p.get("p90", p.get("max"))}
            if all(_finite_number(fields[k]) for k in fields) and fields["p10"] <= fields["p25"] <= fields["p50"] <= fields["p75"] <= fields["p90"]:
                data.append({"label": label, **{k: float(v) for k, v in fields.items()}})
                whisker_definitions.add("p10-p90" if has_percentile_whiskers else "min-max")
        if data:
            charts.append(make_chart(f"distribution_{metric}", "boxplot", f"Distribution of {metric}", "label", "p50", data,
                                     f"Box-and-whisker summary uses saved {', '.join(sorted(whisker_definitions))} whiskers and quartiles; raw observations are not reconstructed.", {"source_method": "persisted_distribution_profile", "whisker_definition": sorted(whisker_definitions)}))
            return charts

    bridge = _as_dict(numbers.get("waterfall_bridge"))
    steps = bridge.get("steps")
    if isinstance(steps, list) and len(steps) >= 2 and re.search(r"change|drop|declin|increase|decreas|shift|impact|period", question, re.I):
        data, previous = [], 0.0
        valid = True
        for index, item in enumerate(steps):
            item = _as_dict(item)
            delta, subtotal = item.get("delta"), item.get("subtotal")
            if not _finite_number(delta) or not _finite_number(subtotal):
                valid = False
                break
            end = float(subtotal)
            data.append({"step": str(item.get("step") or f"Step {index + 1}"), "delta": float(delta), "subtotal": end, "start": 0.0 if index == 0 else previous, "end": end})
            previous = end
        if valid:
            charts.append(make_chart(f"waterfall_{metric}", "waterfall", f"What contributed to the change in {metric}", "step", "delta", data,
                                     "Contribution bars use the saved reconciliation bridge.", {"source_method": "persisted_waterfall_bridge"}))
    return charts


def _build_report_contract(investigation_data: Dict[str, Any]) -> Dict[str, Any]:
    """Build report sections from the saved question, plan, execution record, and output."""
    result = _as_dict(investigation_data.get("analyst_result"))
    numbers = _as_dict(result.get("numbers"))
    plan = _as_dict(investigation_data.get("analysis_plan"))
    contract = _as_dict(numbers.get("canonical_contract") or investigation_data.get("locked_canonical_contract"))
    estimand = _as_dict(plan.get("estimand"))
    semantics = _as_dict(plan.get("semantics"))
    execution = _as_dict(investigation_data.get("execution_record") or numbers.get("execution_record"))
    interpreted: List[str] = []
    calculation: List[str] = []

    kind = str(result.get("kind") or "ANALYSIS").replace("_", " ").title()
    interpreted.append(f"Analysis type: {kind}.")
    metric = execution.get("target_column") or numbers.get("target_column") or contract.get("target_column") or estimand.get("target")
    aggregation = numbers.get("operation") or numbers.get("aggregation") or contract.get("requested_aggregation") or estimand.get("requested_aggregation")
    grouping = execution.get("grouping_columns") or numbers.get("grouping_columns") or contract.get("grouping_columns") or semantics.get("grouping_columns") or []
    comparisons = numbers.get("comparison_values") or contract.get("comparison_values") or semantics.get("comparison_values") or []
    filters_interpreted = numbers.get("filters") or contract.get("filter_conditions") or semantics.get("filters")
    time_scope = contract.get("time_range") or numbers.get("time_window") or numbers.get("period") or semantics.get("time_range")
    if metric:
        interpreted.append(f"Target metric / outcome: {metric}.")
    if aggregation:
        interpreted.append(f"Requested aggregation / method: {aggregation}.")
    if isinstance(grouping, list) and grouping:
        interpreted.append(f"Grouping dimension(s): {', '.join(map(str, grouping))}.")
    if isinstance(comparisons, list) and comparisons:
        interpreted.append(f"Requested comparison values: {', '.join(map(str, comparisons))}.")
    if time_scope:
        interpreted.append(f"Requested time scope: {_report_string(time_scope)}.")
    if filters_interpreted is not None:
        interpreted.append(f"Filters in the interpreted contract: {_report_string(filters_interpreted) if filters_interpreted else 'none recorded'}.")
    datasets = investigation_data.get("dataset_scope")
    if isinstance(datasets, list) and datasets:
        interpreted.append(f"Dataset scope: {', '.join(map(str, datasets))}.")

    if execution:
        actual_metric = execution.get("target_column") or metric
        actual_aggregation = execution.get("aggregation") or aggregation
        actual_groups = execution.get("grouping_columns") or grouping
        if actual_metric or actual_aggregation:
            calculation.append(f"Executed calculation: {actual_aggregation or 'recorded method'}({actual_metric or 'recorded outcome'}).")
        if isinstance(actual_groups, list) and actual_groups:
            calculation.append(f"Grouping actually executed: {', '.join(map(str, actual_groups))}.")
        executed_filters = execution.get("filters")
        if executed_filters is not None:
            calculation.append(f"Filters actually executed: {_report_string(executed_filters) if executed_filters else 'none recorded'}.")
        executed_time = execution.get("time_range")
        if executed_time:
            calculation.append(f"Time window actually executed: {_report_string(executed_time)}.")
        if _finite_number(execution.get("population_row_count")):
            calculation.append(f"Eligible input records: {int(execution['population_row_count']):,}.")
        if _finite_number(execution.get("result_row_count")):
            calculation.append(f"Returned rows / groups: {int(execution['result_row_count']):,}.")
        if execution.get("contract_status") or execution.get("verification_status"):
            calculation.append(f"Execution contract: {execution.get('contract_status', 'not recorded')}; verification: {execution.get('verification_status', 'not recorded')}.")
    else:
        if aggregation or metric:
            calculation.append("No reconciled execution record was attached to this report; method and metric details above describe the saved interpretation, not proof of what ran.")
    method = numbers.get("statistical_test") or numbers.get("method") or _as_dict(numbers.get("trend_inference")).get("method")
    if method:
        calculation.append(f"Recorded statistical method: {method}.")
    if _finite_number(numbers.get("excluded_row_count")) and numbers.get("excluded_row_count", 0) > 0:
        calculation.append(f"Rows excluded by the recorded analysis: {int(numbers['excluded_row_count']):,}.")
    if _finite_number(numbers.get("n")):
        calculation.append(f"Reported sample size: {int(numbers['n']):,}.")

    # Keep the document compact: preserve order while removing duplicates.
    interpreted = list(dict.fromkeys(interpreted))
    calculation = list(dict.fromkeys(calculation))
    visualizations = _select_report_visualizations(investigation_data)
    narrative = compose_report_narrative(
        question=str(investigation_data.get("question") or ""),
        analyst_result=result,
        execution_record=execution,
        evidence=investigation_data.get("evidence") or [],
        status=str(investigation_data.get("status") or ""),
    )
    # The same saved report contract feeds browser and export. Claims are composed from
    # verified facts, and the claim audit is part of the canonical report hash.
    return {
        "schema_version": "1.1",
        "interpreted": interpreted,
        "calculation": calculation,
        "visualizations": visualizations,
        "narrative": narrative,
    }


def _narrative_evidence_label(narrative: Dict[str, Any], claim_text: str) -> str:
    """Return only the verified evidence identifiers attached to this exact claim."""
    claims = narrative.get("claims") if isinstance(narrative.get("claims"), list) else []
    for raw in claims:
        claim = _as_dict(raw)
        if claim.get("text") != claim_text or str(claim.get("validation_status") or "").upper() != "SUPPORTED":
            continue
        refs = [str(ref) for ref in (claim.get("evidence_refs") or []) if str(ref).strip()]
        return "Evidence support: " + ", ".join(refs) if refs else "Evidence support: unavailable"
    return "Evidence support: unavailable"


def _report_string(value: Any) -> str:
    if isinstance(value, str):
        return value
    try:
        return json.dumps(_json_safe(value), sort_keys=True, ensure_ascii=False)
    except (TypeError, ValueError):
        return str(value)


def _report_chart_axis_bounds(ctype: str, vals: List[float]) -> Tuple[float, float]:
    """Choose legible chart bounds without moving non-negative bars below zero."""
    if not vals:
        return 0.0, 1.0
    chart_type = str(ctype or "").lower()
    low_value, high_value = min(vals), max(vals)
    baseline_types = {"bar", "grouped_bar", "stacked_bar", "histogram", "waterfall"}
    if chart_type == "donut":
        return low_value, high_value  # Donut charts have no numeric axes.
    if chart_type in baseline_types and low_value >= 0:
        low = 0.0
        high = high_value
        if high <= 0:
            high = 1.0
        high += (high - low) * 0.06
        return low, high
    if chart_type in baseline_types and high_value <= 0:
        low = low_value
        high = 0.0
        if low >= 0:
            low = -1.0
        low -= (high - low) * 0.06
        return low, high
    low, high = low_value, high_value
    if high == low:
        high = low + 1.0
    pad = (high - low) * 0.06
    return low - pad, high + pad


def _render_report_chart_png(chart: Dict[str, Any], width: int = 1000, height: int = 440) -> Optional[bytes]:
    """Draw a compact report chart using Pillow, without a plotting-runtime dependency."""
    import math

    rows = chart.get("data") if isinstance(chart.get("data"), list) else []
    if not rows:
        return None
    ctype = chart.get("chart_type")
    x_field = _as_dict(chart.get("x_axis")).get("field") or "category"
    y_field = _as_dict(chart.get("y_axis")).get("field") or "value"
    if ctype not in {"line", "multi_line", "scatter", "bar", "grouped_bar", "stacked_bar", "histogram", "donut", "boxplot", "waterfall"}:
        return None
    if ctype in {"line", "multi_line", "scatter"}:
        data_points = [(str(row.get(x_field, "")), float(row[y_field])) for row in rows if isinstance(row, dict) and _finite_number(row.get(y_field))]
        if not data_points:
            return None
        groups: Dict[str, List[Tuple[str, float]]] = {}
        sf = chart.get("series_field")
        for row in rows:
            if not isinstance(row, dict) or not _finite_number(row.get(y_field)):
                continue
            group = str(row.get(sf) or "Series 1") if sf else "Series 1"
            groups.setdefault(group, []).append((str(row.get(x_field, "")), float(row[y_field])))
        vals = [v for items in groups.values() for _, v in items]
        # Include interval endpoints in the report's scale and draw them when the
        # source saved confidence/prediction intervals with each observation.
        md = _as_dict(chart.get("metadata"))
        ci_low_field = str(md.get("ci_low_field") or "ci_low")
        ci_high_field = str(md.get("ci_high_field") or "ci_high")
        interval_rows = [row for row in rows if isinstance(row, dict)
                         and _finite_number(row.get(ci_low_field)) and _finite_number(row.get(ci_high_field))
                         and float(row[ci_low_field]) <= float(row[ci_high_field])]
        vals.extend(float(row[field]) for row in interval_rows for field in (ci_low_field, ci_high_field))
    elif ctype == "boxplot":
        vals = [float(row[k]) for row in rows if isinstance(row, dict) for k in ("p10", "p25", "p50", "p75", "p90") if _finite_number(row.get(k))]
        groups = {}
        data_points = []
    elif ctype == "waterfall":
        vals = [float(row.get("start", 0)) for row in rows if isinstance(row, dict) and _finite_number(row.get("start"))] + [float(row.get("end", row.get("delta"))) for row in rows if isinstance(row, dict) and _finite_number(row.get("end", row.get("delta")))]
        groups = {}
        data_points = []
    elif ctype in {"bar", "grouped_bar", "stacked_bar", "histogram"}:
        data_points = [(str(row.get(x_field, "")), float(row[y_field])) for row in rows if isinstance(row, dict) and _finite_number(row.get(y_field))]
        if not data_points:
            return None
        vals = [v for _, v in data_points]
        sf = chart.get("series_field")
        groups = {}
        for row in rows:
            if not isinstance(row, dict) or not _finite_number(row.get(y_field)):
                continue
            group = str(row.get(sf) or "Series 1") if sf else "Series 1"
            groups.setdefault(group, []).append((str(row.get(x_field, "")), float(row[y_field])))
        if ctype == "stacked_bar":
            positive: Dict[str, float] = {}
            negative: Dict[str, float] = {}
            for category, value in data_points:
                if value >= 0:
                    positive[category] = positive.get(category, 0.0) + value
                else:
                    negative[category] = negative.get(category, 0.0) + value
            vals.extend(positive.values())
            vals.extend(negative.values())
    else:
        vals = [float(row.get(y_field)) for row in rows if isinstance(row, dict) and _finite_number(row.get(y_field))]
        groups = {}
        data_points = []
    if not vals:
        return None
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    ink = (30, 41, 59)
    muted = (100, 116, 139)
    grid = (226, 232, 240)
    palette = [(8, 145, 178), (124, 58, 237), (22, 163, 74), (217, 119, 6), (225, 29, 72), (37, 99, 235)]
    title = str(chart.get("title") or "Visualization")
    draw.text((48, 20), title[:110], fill=ink)
    explanation = str(chart.get("explanation") or "")
    if explanation:
        draw.text((48, 47), explanation[:145], fill=muted)
    left, top, right, bottom = 92, 88, width - 38, height - 70
    plot_w, plot_h = right - left, bottom - top
    # Non-negative bars start at zero; negative-axis padding on positive-only data
    # previously suggested impossible values in exported report charts. Donuts have no axes.
    low, high = _report_chart_axis_bounds(str(ctype), vals)
    def y_pos(v: float) -> float:
        return bottom - ((v - low) / (high - low)) * plot_h
    if ctype != "donut":
        for tick in range(5):
            val = low + (high - low) * tick / 4
            y = y_pos(val)
            draw.line((left, y, right, y), fill=grid, width=1)
            draw.text((6, y - 7), _fmt_compact(val), fill=muted)
        draw.line((left, top, left, bottom), fill=muted, width=2)
        draw.line((left, bottom, right, bottom), fill=muted, width=2)

    if ctype == "donut":
        entries = [(str(row.get(x_field, "")), float(row.get(y_field))) for row in rows if isinstance(row, dict) and _finite_number(row.get(y_field)) and float(row.get(y_field)) >= 0]
        total = sum(v for _, v in entries)
        if not entries or total <= 0:
            return None
        # Draw the ring in the left half with a compact legend on the right.
        box = (left + 10, top + 5, left + min(plot_h - 5, plot_w * 0.52), bottom - 5)
        diameter = min(box[2] - box[0], box[3] - box[1])
        box = (box[0], box[1], box[0] + diameter, box[1] + diameter)
        start = -90.0
        for i, (name, value) in enumerate(entries):
            end = start + 360.0 * value / total
            draw.pieslice(box, start=start, end=end, fill=palette[i % len(palette)])
            start = end
        inner = diameter * 0.58
        cx, cy = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
        draw.ellipse((cx-inner/2, cy-inner/2, cx+inner/2, cy+inner/2), fill="white")
        draw.text((cx-inner/2+8, cy-5), _fmt_compact(total), fill=ink)
        lx, ly = box[2] + 25, top + 15
        for i, (name, value) in enumerate(entries[:12]):
            draw.rectangle((lx, ly+3, lx+10, ly+13), fill=palette[i % len(palette)])
            draw.text((lx+17, ly), f"{name[:27]}  {_fmt_compact(value)} ({value/total:.1%})", fill=ink)
            ly += 25
    elif ctype == "boxplot":
        boxes = [row for row in rows if isinstance(row, dict) and all(_finite_number(row.get(k)) for k in ("p10", "p25", "p50", "p75", "p90"))]
        if not boxes:
            return None
        band = plot_w / len(boxes)
        for index, row in enumerate(boxes):
            cx = left + band * (index + 0.5)
            y10, y25, y50, y75, y90 = (y_pos(float(row[k])) for k in ("p10", "p25", "p50", "p75", "p90"))
            col = palette[index % len(palette)]
            draw.line((cx, y90, cx, y10), fill=col, width=3)
            draw.line((cx-16, y90, cx+16, y90), fill=col, width=3)
            draw.line((cx-16, y10, cx+16, y10), fill=col, width=3)
            draw.rectangle((cx-32, y75, cx+32, y25), outline=col, fill=(232, 240, 250), width=3)
            draw.line((cx-32, y50, cx+32, y50), fill=col, width=4)
            draw.text((cx-48, bottom+14), str(row.get("label") or f"Group {index+1}")[:14], fill=muted)
    elif ctype == "waterfall":
        n = len(rows)
        if not n:
            return None
        band = plot_w / n
        for index, row in enumerate(rows):
            if not isinstance(row, dict) or not _finite_number(row.get("delta")):
                continue
            x = left + band * (index + 0.5)
            start_v = float(row.get("start", 0)) if _finite_number(row.get("start", 0)) else 0.0
            end_v = float(row.get("end", row.get("delta"))) if _finite_number(row.get("end", row.get("delta"))) else float(row["delta"])
            y1, y2 = y_pos(start_v), y_pos(end_v)
            col = palette[2] if float(row["delta"]) >= 0 else palette[4]
            draw.rectangle((x-band*0.28, min(y1,y2), x+band*0.28, max(y1,y2)+1), fill=col)
            draw.text((x-band*0.45, bottom+14), str(row.get("step") or index+1)[:12], fill=muted)
    elif ctype in {"bar", "grouped_bar", "stacked_bar", "histogram"}:
        categories = list(dict.fromkeys(name for name, _ in data_points))
        series_values: Dict[str, Dict[str, float]] = {}
        sf = chart.get("series_field")
        for row in rows:
            if not isinstance(row, dict) or not _finite_number(row.get(y_field)):
                continue
            cat = str(row.get(x_field, ""))
            series = str(row.get(sf) or "Series 1") if sf else "Series 1"
            series_values.setdefault(series, {})[cat] = float(row[y_field])
        band = plot_w / max(1, len(categories))
        total_series = max(1, len(series_values))
        bar_w = max(2, band * (0.94 if ctype == "histogram" else 0.68) / (1 if ctype in {"stacked_bar", "histogram"} else total_series))
        y_zero = y_pos(0)
        for ci, cat in enumerate(categories):
            if ctype == "stacked_bar":
                positive_base = 0.0
                negative_base = 0.0
                x = left + band * ci + (band - bar_w) / 2
                for si, (series, values) in enumerate(series_values.items()):
                    if cat not in values:
                        continue
                    value = values[cat]
                    if value >= 0:
                        start_v, end_v = positive_base, positive_base + value
                        positive_base = end_v
                    else:
                        start_v, end_v = negative_base, negative_base + value
                        negative_base = end_v
                    y1, y2 = y_pos(start_v), y_pos(end_v)
                    draw.rectangle((x, min(y1, y2), x+bar_w-2, max(y1, y2)+1), fill=palette[si % len(palette)])
            else:
                for si, (series, values) in enumerate(series_values.items()):
                    if cat not in values:
                        continue
                    x = left + band * ci + band*0.16 + si*bar_w
                    yy = y_pos(values[cat])
                    draw.rectangle((x, min(yy,y_zero), x+bar_w-2, max(yy,y_zero)+1), fill=palette[si % len(palette)])
            draw.text((left+band*ci+2, bottom+14), cat[:13], fill=muted)
    elif ctype in {"line", "multi_line", "scatter"}:
        # Sort by x string for dates/ISO timestamps and by numeric value when all x's parse.
        categories = list(dict.fromkeys(x for x, _ in data_points))
        try:
            categories.sort(key=float)
        except (ValueError, TypeError):
            categories.sort()
        x_at = {x: left + i * plot_w / max(1, len(categories)-1) for i, x in enumerate(categories)}
        if ctype in {"line", "multi_line"}:
            md = _as_dict(chart.get("metadata"))
            ci_low_field = str(md.get("ci_low_field") or "ci_low")
            ci_high_field = str(md.get("ci_high_field") or "ci_high")
            for row in rows:
                if not isinstance(row, dict) or not _finite_number(row.get(ci_low_field)) or not _finite_number(row.get(ci_high_field)):
                    continue
                if float(row[ci_low_field]) > float(row[ci_high_field]):
                    continue
                x = x_at.get(str(row.get(x_field, "")))
                if x is None:
                    continue
                y_low, y_high = y_pos(float(row[ci_low_field])), y_pos(float(row[ci_high_field]))
                draw.line((x, y_high, x, y_low), fill=muted, width=2)
                draw.line((x-5, y_high, x+5, y_high), fill=muted, width=2)
                draw.line((x-5, y_low, x+5, y_low), fill=muted, width=2)
        for si, (series, points) in enumerate(groups.items()):
            try:
                points = sorted(points, key=lambda item: float(item[0]))
            except (ValueError, TypeError):
                points = sorted(points, key=lambda item: item[0])
            coords = [(x_at[x], y_pos(val)) for x, val in points if x in x_at]
            col = palette[si % len(palette)]
            if ctype != "scatter" and len(coords) >= 2:
                draw.line(coords, fill=col, width=3)
            for x, y in coords:
                draw.ellipse((x-4, y-4, x+4, y+4), fill=col)
        for idx, x in enumerate(categories):
            if len(categories) <= 12 or idx in {0, len(categories)-1} or idx % max(1, len(categories)//10) == 0:
                xpos = x_at[x]
                draw.text((xpos-22, bottom+14), x[:12], fill=muted)
    out = BytesIO()
    img.save(out, format="PNG", optimize=True)
    return out.getvalue()


def _report_is_verified(investigation_data: Dict[str, Any]) -> bool:
    if str(investigation_data.get("status") or "").upper() != "COMPLETED":
        return False
    result = _as_dict(investigation_data.get("analyst_result"))
    numbers = _as_dict(result.get("numbers"))
    execution = _as_dict(investigation_data.get("execution_record") or numbers.get("execution_record"))
    if execution:
        return str(execution.get("contract_status") or "").upper() == "MATCHED" and str(execution.get("verification_status") or "").upper() == "VERIFIED"
    evidence = investigation_data.get("evidence") or []
    return bool(evidence) and all(str(_as_dict(item).get("validation_status") or "").upper() in {"VERIFIED", "PASSED"} for item in evidence)

def canonical_report_payload(inv_data: Dict[str, Any]) -> Dict[str, Any]:
    """Extract a canonical, normalized dictionary representation of the report content.

    Any material change in the investigation identity, question, headline, metrics,
    tables, insights, limitations, recommended actions, manifest hash, or evidence
    will alter this canonical payload.
    """
    analyst_result = inv_data.get("analyst_result") or {}
    numbers = analyst_result.get("numbers") or {}
    evidence_items = inv_data.get("evidence") or []

    # Sort evidence items deterministically by ID or statement
    sorted_evidence = []
    for ev in evidence_items:
        sorted_evidence.append({
            "id": str(ev.get("id", "")),
            "statement": str(ev.get("statement", "")),
            "calculation_summary": str(ev.get("calculation_summary", "")),
            "row_count_analyzed": ev.get("row_count_analyzed"),
            "validation_status": str(ev.get("validation_status", "UNVERIFIED")),
            "sql_executed": str(ev.get("sql_executed") or ""),
        })
    sorted_evidence.sort(key=lambda x: (x["id"], x["statement"]))

    return {
        "investigation_id": str(inv_data.get("id", "")),
        "question": str(inv_data.get("question", "")).strip(),
        "status": str(inv_data.get("status", "")).strip().upper(),
        "verdict": str(inv_data.get("verdict_type") or inv_data.get("verdict") or "").strip().upper(),
        "headline": str(analyst_result.get("headline") or inv_data.get("headline") or inv_data.get("direct_answer") or "").strip(),
        "kind": str(analyst_result.get("kind") or "ANALYSIS").strip().upper(),
        "numbers": {
            k: v for k, v in sorted(numbers.items(), key=lambda item: item[0])
            if isinstance(v, (str, int, float, bool, dict, list))
        },
        "insights": [str(d).strip() for d in (analyst_result.get("details") or []) if str(d).strip()][:3],
        "limitations": [str(c).strip() for c in (analyst_result.get("caveats") or []) if str(c).strip()],
        "actions": [str(a).strip() for a in (analyst_result.get("next_steps") or []) if str(a).strip()],
        "manifest_hash": str(inv_data.get("reproducible_manifest_hash") or "").strip(),
        # The report's interpretation and execution facts are part of the sealed payload.
        # A changed scope, aggregation, execution status, or visualization must change the seal.
        "analysis_plan": _json_safe(_as_dict(inv_data.get("analysis_plan"))),
        "execution_record": _json_safe(_as_dict(inv_data.get("execution_record") or numbers.get("execution_record"))),
        "report_contract": _json_safe(_build_report_contract(inv_data)),
        "evidence": sorted_evidence,
    }


def compute_canonical_report_hash(inv_data: Dict[str, Any]) -> str:
    """Compute full 64-character SHA-256 hash of the canonical report payload."""
    payload = canonical_report_payload(inv_data)
    canonical_json = json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=True)
    return hashlib.sha256(canonical_json.encode("utf-8")).hexdigest()


def compute_export_fingerprint(inv_data: Dict[str, Any]) -> str:
    """Compute the user-facing short document seal from the canonical full report hash.

    Format: MDED-XXXX-XXXX-XXXX-XXXX (first 16 hex characters).
    Backed by the 64-character canonical report hash for complete tamper-evident audit.
    """
    full_hash = compute_canonical_report_hash(inv_data)
    return f"MDED-{full_hash[:4].upper()}-{full_hash[4:8].upper()}-{full_hash[8:12].upper()}-{full_hash[12:16].upper()}"


def verify_report_integrity(
    inv_data: Dict[str, Any],
    expected_full_hash: Optional[str] = None,
    expected_seal: Optional[str] = None,
) -> Tuple[bool, Dict[str, Any]]:
    """Verify whether a report's data matches an expected full hash and/or visible seal.

    Returns:
        (is_valid, details_dict)
    """
    computed_full_hash = compute_canonical_report_hash(inv_data)
    computed_seal = f"MDED-{computed_full_hash[:4].upper()}-{computed_full_hash[4:8].upper()}-{computed_full_hash[8:12].upper()}-{computed_full_hash[12:16].upper()}"

    matches_full = (expected_full_hash.lower() == computed_full_hash.lower()) if expected_full_hash else True
    matches_seal = (expected_seal.upper() == computed_seal.upper()) if expected_seal else True

    is_valid = matches_full and matches_seal
    return is_valid, {
        "is_valid": is_valid,
        "computed_full_hash": computed_full_hash,
        "computed_seal": computed_seal,
        "matches_full_hash": matches_full,
        "matches_seal": matches_seal,
    }


def _generate_minded_logo_png(size: int = 120) -> bytes:
    """Draw a clean, crisp MindEd hexagonal lattice prism emblem as PNG bytes."""
    img = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)

    # Coordinates for hexagon
    cx, cy = size / 2, size / 2
    r = size * 0.44

    import math
    points = []
    for i in range(6):
        angle = math.radians(60 * i - 30)
        px = cx + r * math.cos(angle)
        py = cy + r * math.sin(angle)
        points.append((px, py))

    # Dark graphite backing
    draw.polygon(points, fill=(17, 24, 39, 255), outline=(245, 166, 35, 255), width=max(2, int(size * 0.05)))

    # Amber interior lattice lines
    amber_inner = (255, 204, 112, 230)
    w_line = max(1, int(size * 0.035))
    draw.line([points[0], points[3]], fill=amber_inner, width=w_line)
    draw.line([points[1], points[4]], fill=amber_inner, width=w_line)
    draw.line([points[2], points[5]], fill=amber_inner, width=w_line)

    # Core center circle (representing nucleus of intelligence)
    r_core = size * 0.12
    draw.ellipse(
        [cx - r_core, cy - r_core, cx + r_core, cy + r_core],
        fill=(245, 166, 35, 255),
        outline=(255, 255, 255, 200),
        width=1,
    )

    out = BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out.getvalue()


def _question_requests_ranking(question: str) -> bool:
    return bool(re.search(r"\b(?:highest|lowest|most|least|largest|smallest|top|bottom|best|worst|rank(?:ed|ing)?)\b", str(question or ""), re.I))


def _extract_table_data(kind: str, numbers: Dict[str, Any], question: str = "") -> Optional[Dict[str, Any]]:
    kind = (kind or "").upper()
    raw_table = numbers.get("table")
    question = str(question or "")
    stored_contract = _as_dict(numbers.get("canonical_contract"))
    stored_direction = str(stored_contract.get("ranking_direction") or "").upper()
    is_ranking = kind == "RANKING" or _question_requests_ranking(question)
    if is_ranking and isinstance(raw_table, dict):
        entries = []
        for name, data in raw_table.items():
            val = data.get("value", 0.0) if isinstance(data, dict) else data
            n_rows = data.get("n", 0) if isinstance(data, dict) else 0
            if not _finite_number(val):
                continue
            entries.append({"name": str(name), "value": float(val), "rows": int(n_rows) if _finite_number(n_rows) else 0})
        if not entries:
            return None
        low_intent = bool(re.search(r"\b(?:lowest|least|smallest|worst|bottom)\b", question, re.I))
        high_intent = bool(re.search(r"\b(?:highest|most|largest|best|top)\b", question, re.I))
        if low_intent and not high_intent:
            ascending = True
        elif high_intent and not low_intent:
            ascending = False
        elif stored_direction in {"ASC", "DESC"}:
            ascending = stored_direction == "ASC"
        else:
            ascending = False
        entries.sort(key=lambda item: item["value"], reverse=not ascending)

        operation = str(numbers.get("operation") or numbers.get("aggregation") or stored_contract.get("requested_aggregation") or "").upper()
        additive = operation in {"SUM", "TOTAL", "COUNT", "ROW_COUNT"}
        grand = sum(item["value"] for item in entries)
        include_share = additive and grand > 0 and all(item["value"] >= 0 for item in entries)
        if additive:
            value_header = "Total Value"
        elif operation in {"MEAN", "AVG", "AVERAGE"}:
            value_header = "Average Value"
        else:
            value_header = "Value"
        headers = ["Rank", "Group / Entity", value_header, "Rows"]
        if include_share:
            headers.append("Share")
        rows = []
        for idx, item in enumerate(entries, start=1):
            row = [str(idx), item["name"], _fmt_compact(item["value"]), f"{item['rows']:,}"]
            if include_share:
                share_str = f"{(item['value'] / grand * 100):.1f}%"
                row.append(share_str)
            rows.append(row)
        return {"title": "Ranking by Metric", "headers": headers, "rows": rows}

    if kind in ("ASSOCIATION", "CORRELATION", "REGRESSION"):
        rows = []
        if "pearson_r" in numbers and numbers["pearson_r"] is not None:
            rows.append(["Pearson correlation (r)", f"{float(numbers['pearson_r']):.3f}"])
        if "ci_low" in numbers and "ci_high" in numbers:
            rows.append(["95% confidence interval", f"{float(numbers['ci_low']):.3f} to {float(numbers['ci_high']):.3f}"])
        if "pearson_p" in numbers and numbers["pearson_p"] is not None:
            p = float(numbers["pearson_p"])
            rows.append(["P-value", "p < 0.001" if p < 0.001 else f"{p:.4f}"])
        if "slope" in numbers and numbers["slope"] is not None:
            slope = float(numbers["slope"])
            rows.append(["Simple regression slope", f"{slope:+.4f}"])
        if "r_squared" in numbers and numbers["r_squared"] is not None:
            rows.append(["Variation explained (R²)", f"{(float(numbers['r_squared']) * 100):.2f}%"])

        multiv = numbers.get("multivariate_context")
        if isinstance(multiv, dict):
            if "controlled_slope" in multiv:
                rows.append(["Adjusted slope (controlled)", f"{float(multiv['controlled_slope']):+.4f}"])
            if "adjusted_r_squared" in multiv:
                rows.append(["Adjusted model R²", f"{(float(multiv['adjusted_r_squared']) * 100):.2f}%"])

        if rows:
            return {"title": "Statistical Evidence", "headers": ["Metric", "Result"], "rows": rows}

    if isinstance(raw_table, dict):
        rows = []
        for k, v in list(raw_table.items())[:30]:
            val_str = _fmt_compact(v.get("value")) if isinstance(v, dict) else str(v)
            rows.append([str(k), val_str])
        return {"title": "Results Table", "headers": ["Item / Group", "Value"], "rows": rows}

    raw_list = numbers.get("table") or numbers.get("rows")
    if isinstance(raw_list, list) and raw_list and isinstance(raw_list[0], dict):
        headers = list(raw_list[0].keys())
        rows = []
        for row_dict in raw_list[:40]:
            row = []
            for h in headers:
                val = row_dict.get(h)
                if isinstance(val, float):
                    row.append(_fmt_compact(val))
                elif isinstance(val, int):
                    row.append(f"{val:,}")
                else:
                    row.append(str(val) if val is not None else "")
            rows.append(row)
        return {"title": "Results Table", "headers": [str(h).replace("_", " ").title() for h in headers], "rows": rows}

    return None


def export_investigation_to_pdf(investigation_data: Dict[str, Any]) -> bytes:
    """Generate professional publication-quality PDF matching MindEd design standard."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )
    styles = getSampleStyleSheet()

    brand_style = ParagraphStyle(
        "BrandHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#B45309"),
    )
    fingerprint_style = ParagraphStyle(
        "Fingerprint",
        parent=styles["Normal"],
        fontName="Courier-Bold",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#475569"),
        alignment=2,
    )
    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#111827"),
    )
    headline_style = ParagraphStyle(
        "Headline",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=17,
        textColor=colors.HexColor("#1E293B"),
    )
    section_h_style = ParagraphStyle(
        "SectionHeader",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor("#475569"),
        textTransform="uppercase",
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#334155"),
    )
    callout_style = ParagraphStyle(
        "Callout",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor("#92400E"),
    )
    metric_val_style = ParagraphStyle(
        "MetricVal",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#D97706"),
    )
    sub_metric_style = ParagraphStyle(
        "SubMetric",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#64748B"),
    )

    story = []

    # 1. Header with Emblem Logo + Fingerprint Banner
    logo_bytes = _generate_minded_logo_png(size=80)
    logo_img = RLImage(BytesIO(logo_bytes), width=0.42 * inch, height=0.42 * inch)
    fingerprint_code = compute_export_fingerprint(investigation_data)
    canonical_hash = compute_canonical_report_hash(investigation_data)

    header_table_data = [
        [
            logo_img,
            Paragraph("<b>MINDED AA-OS</b><br/><font color='#64748B' size=8>Autonomous Analytical Operating System</font>", brand_style),
            Paragraph(f"<b>AUTHENTIC EXPORT SEAL</b><br/><font color='#D97706'><b>{fingerprint_code}</b></font>", fingerprint_style),
        ]
    ]
    header_table = Table(header_table_data, colWidths=[0.55 * inch, 3.6 * inch, 2.85 * inch])
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#F5A623"), spaceAfter=10))

    # Question Title
    question = investigation_data.get("question", "Analytical Investigation")
    story.append(Paragraph(html.escape(str(question)), title_style))
    story.append(Spacer(1, 4))

    verdict = investigation_data.get("verdict_type") or investigation_data.get("verdict") or "OBSERVED"
    status = investigation_data.get("status", "COMPLETED")
    date_str = datetime.now(timezone.utc).strftime("%b %d, %Y · %H:%M UTC")
    meta_text = f"<b>Status:</b> {html.escape(str(status))} &nbsp;|&nbsp; <b>Verdict:</b> {html.escape(str(verdict))} &nbsp;|&nbsp; <b>Certified:</b> {date_str}"
    story.append(Paragraph(meta_text, sub_metric_style))
    story.append(Spacer(1, 10))

    # Analyst Result Extraction
    analyst_result = investigation_data.get("analyst_result") or {}
    headline = analyst_result.get("headline") or investigation_data.get("headline") or investigation_data.get("direct_answer") or ""
    numbers = analyst_result.get("numbers") or {}
    kind = analyst_result.get("kind") or "ANALYSIS"
    report_contract = _build_report_contract(investigation_data)

    # 2. Headline
    if headline:
        story.append(Paragraph("HEADLINE", section_h_style))
        story.append(Spacer(1, 3))
        story.append(Paragraph(html.escape(str(headline)), headline_style))
        story.append(Spacer(1, 10))

    narrative = _as_dict(report_contract.get("narrative"))
    if narrative.get("status") in {"READY", "PARTIAL", "LIMITED"} and narrative.get("summary"):
        story.append(Paragraph(html.escape(str(narrative.get("title") or "Analysis explanation")), section_h_style))
        story.append(Spacer(1, 3))
        summary_text = str(narrative["summary"])
        story.append(Paragraph(html.escape(summary_text), body_style))
        story.append(Paragraph(html.escape(_narrative_evidence_label(narrative, summary_text)), sub_metric_style))
        for paragraph in narrative.get("paragraphs") or []:
            paragraph_text = str(paragraph)
            story.append(Spacer(1, 3))
            story.append(Paragraph(html.escape(paragraph_text), body_style))
            story.append(Paragraph(html.escape(_narrative_evidence_label(narrative, paragraph_text)), sub_metric_style))
        supported_count = narrative.get("supported_claim_count")
        if _finite_number(supported_count):
            story.append(Spacer(1, 3))
            story.append(Paragraph(html.escape(f"Evidence-linked claims: {int(supported_count)} supported; unsupported claims are withheld."), sub_metric_style))
        story.append(Spacer(1, 6))

    # 3. Key Result Metric Banner
    top_val = numbers.get("top_value")
    pearson_r = numbers.get("pearson_r")
    if top_val is not None or pearson_r is not None:
        val_str = _fmt_compact(top_val) if top_val is not None else f"r = {float(pearson_r):.3f}"
        sub_str = f"Category: {kind} · Sample: {numbers.get('n', 'N/A')} observations"
        story.append(Paragraph(val_str, metric_val_style))
        story.append(Paragraph(sub_str, sub_metric_style))
        story.append(Spacer(1, 10))

    # 4. Structured Results Table
    table_meta = _extract_table_data(kind, numbers, str(investigation_data.get("question") or ""))
    if table_meta:
        story.append(Paragraph(table_meta["title"], section_h_style))
        story.append(Spacer(1, 4))
        table_content = [table_meta["headers"]] + table_meta["rows"]
        t = Table(table_content, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F8FAFC")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#334155")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
            ("TOPPADDING", (0, 0), (-1, -1), 4.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 1), (-1, -1), 9),
            ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#1E293B")),
        ]))
        story.append(t)
        story.append(Spacer(1, 10))

    # 5. Interpretation, calculation transparency, and visual evidence
    if report_contract.get("interpreted"):
        story.append(Paragraph("WHAT MINDED INTERPRETED", section_h_style))
        story.append(Spacer(1, 4))
        for item in report_contract["interpreted"]:
            story.append(Paragraph(f"• {html.escape(str(item))}", body_style))
            story.append(Spacer(1, 2))
        story.append(Spacer(1, 6))

    if report_contract.get("calculation"):
        story.append(Paragraph("HOW THE RESULT WAS CALCULATED", section_h_style))
        story.append(Spacer(1, 4))
        for item in report_contract["calculation"]:
            story.append(Paragraph(f"• {html.escape(str(item))}", body_style))
            story.append(Spacer(1, 2))
        story.append(Spacer(1, 6))

    report_charts = report_contract.get("visualizations") or []
    if report_charts:
        story.append(Paragraph("VISUALIZATIONS", section_h_style))
        story.append(Spacer(1, 4))
        for chart in report_charts[:4]:
            story.append(Paragraph(html.escape(str(chart.get("title") or "Visualization")), headline_style))
            if chart.get("explanation"):
                story.append(Paragraph(html.escape(str(chart["explanation"])), sub_metric_style))
                story.append(Spacer(1, 3))
            png = _render_report_chart_png(chart)
            if png:
                story.append(RLImage(BytesIO(png), width=6.25 * inch, height=2.75 * inch))
                story.append(Spacer(1, 8))
        story.append(Spacer(1, 3))

    # 6. Key Insights
    details = analyst_result.get("details") or []
    if details:
        story.append(Paragraph("KEY INSIGHTS", section_h_style))
        story.append(Spacer(1, 4))
        for d in details[:3]:
            story.append(Paragraph(f"• {html.escape(str(d))}", body_style))
            story.append(Spacer(1, 2.5))
        story.append(Spacer(1, 8))

    # 7. Limitations & Assumptions
    caveats = analyst_result.get("caveats") or []
    narrative_caveats = _as_dict(report_contract.get("narrative")).get("limitations") or []
    caveats = list(dict.fromkeys([*caveats, *narrative_caveats]))
    if caveats or kind in ("ASSOCIATION", "CORRELATION"):
        story.append(Paragraph("LIMITATIONS & ASSUMPTIONS", section_h_style))
        story.append(Spacer(1, 4))
        all_caveats = list(caveats)
        if kind in ("ASSOCIATION", "CORRELATION") and not any("causation" in c.lower() for c in all_caveats):
            all_caveats.append("The analysis establishes statistical association, not mechanistic causation.")
        for c in all_caveats:
            story.append(Paragraph(f"• {html.escape(str(c))}", callout_style))
            story.append(Spacer(1, 2.5))
        story.append(Spacer(1, 8))

    # 8. Recommended Actions
    next_steps = analyst_result.get("next_steps") or []
    if next_steps:
        story.append(Paragraph("RECOMMENDED ACTIONS", section_h_style))
        story.append(Spacer(1, 4))
        for idx, act in enumerate(next_steps, start=1):
            story.append(Paragraph(f"{idx}. {html.escape(str(act))}", body_style))
            story.append(Spacer(1, 2.5))
        story.append(Spacer(1, 8))

    # 9. Evidence & Provenance Audit with Security Seal
    evidence = investigation_data.get("evidence") or []
    manifest_hash = investigation_data.get("reproducible_manifest_hash") or "N/A"
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#E2E8F0"), spaceAfter=8))
    story.append(Paragraph("EVIDENCE AUDIT & PROVENANCE", section_h_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(f"<b>Full Canonical Digest (SHA-256):</b> <font face='Courier' size=7.5>{canonical_hash}</font>", sub_metric_style))
    story.append(Paragraph(f"<b>Manifest Hash:</b> <font face='Courier' size=8>{manifest_hash}</font> &nbsp;|&nbsp; <b>Seal:</b> <font face='Courier' color='#D97706'><b>{fingerprint_code}</b></font>", sub_metric_style))
    story.append(Spacer(1, 4))

    for ev in evidence[:5]:
        stmt = html.escape(str(ev.get("statement", "Evidence Item")))
        v_status = html.escape(str(ev.get("validation_status", "UNVERIFIED")))
        calc = html.escape(str(ev.get("calculation_summary", "")))
        rows_n = ev.get("row_count_analyzed")
        rows_str = f" | Rows analyzed: {rows_n:,}" if rows_n else ""
        story.append(Paragraph(f"• <b>{stmt}</b> [{v_status}]{rows_str}", body_style))
        if calc:
            story.append(Paragraph(f"&nbsp;&nbsp;&nbsp;&nbsp;<i>Calculation:</i> {calc}", sub_metric_style))
        story.append(Spacer(1, 2))


    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def export_investigation_to_docx(investigation_data: Dict[str, Any]) -> bytes:
    """Generate clean structured Microsoft Word document (.docx) with MindEd branding and fingerprint."""
    doc = docx.Document()

    fingerprint_code = compute_export_fingerprint(investigation_data)
    canonical_hash = compute_canonical_report_hash(investigation_data)
    logo_bytes = _generate_minded_logo_png(size=64)

    # Document Header Table (Emblem / Logo + Brand + Fingerprint)
    tbl_hdr = doc.add_table(rows=1, cols=2)
    tbl_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left, c_right = tbl_hdr.rows[0].cells

    p_logo = c_left.paragraphs[0]
    run_img = p_logo.add_run()
    run_img.add_picture(BytesIO(logo_bytes), width=Inches(0.45), height=Inches(0.45))
    run_brand = p_logo.add_run("  MINDED AA-OS\n  Autonomous Analytical Intelligence")
    run_brand.font.size = Pt(9.5)
    run_brand.font.color.rgb = RGBColor(180, 83, 9)
    run_brand.bold = True

    p_fp = c_right.paragraphs[0]
    p_fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_fp_label = p_fp.add_run("AUTHENTIC EXPORT SEAL\n")
    run_fp_label.font.size = Pt(8)
    run_fp_label.font.color.rgb = RGBColor(100, 116, 139)
    run_fp_val = p_fp.add_run(fingerprint_code)
    run_fp_val.font.size = Pt(10)
    run_fp_val.font.color.rgb = RGBColor(217, 119, 6)
    run_fp_val.bold = True

    doc.add_paragraph()

    # Inquiry Title
    h1 = doc.add_heading(level=1)
    run_title = h1.add_run(investigation_data.get("question", "Analytical Investigation Report"))
    run_title.font.size = Pt(17)
    run_title.bold = True

    # Metadata
    verdict = investigation_data.get("verdict_type") or investigation_data.get("verdict") or "OBSERVED"
    status = investigation_data.get("status", "COMPLETED")
    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Status: {status}  |  Verdict: {verdict}  |  Certified: {datetime.now(timezone.utc).strftime('%b %d, %Y UTC')}").italic = True

    analyst_result = investigation_data.get("analyst_result") or {}
    headline = analyst_result.get("headline") or investigation_data.get("headline") or investigation_data.get("direct_answer") or ""
    numbers = analyst_result.get("numbers") or {}
    kind = analyst_result.get("kind") or "ANALYSIS"
    report_contract = _build_report_contract(investigation_data)

    # Headline
    if headline:
        doc.add_heading("Headline", level=2)
        p_head = doc.add_paragraph()
        run_h = p_head.add_run(headline)
        run_h.bold = True
        run_h.font.size = Pt(12)

    narrative = _as_dict(report_contract.get("narrative"))
    if narrative.get("status") in {"READY", "PARTIAL", "LIMITED"} and narrative.get("summary"):
        doc.add_heading(str(narrative.get("title") or "Analysis explanation"), level=2)
        summary_text = str(narrative["summary"])
        doc.add_paragraph(summary_text)
        summary_source = doc.add_paragraph(_narrative_evidence_label(narrative, summary_text))
        summary_source.runs[0].italic = True
        summary_source.runs[0].font.size = Pt(8)
        for paragraph in narrative.get("paragraphs") or []:
            paragraph_text = str(paragraph)
            doc.add_paragraph(paragraph_text)
            source = doc.add_paragraph(_narrative_evidence_label(narrative, paragraph_text))
            source.runs[0].italic = True
            source.runs[0].font.size = Pt(8)
        supported_count = narrative.get("supported_claim_count")
        if _finite_number(supported_count):
            note = doc.add_paragraph()
            run = note.add_run(f"Evidence-linked claims: {int(supported_count)} supported; unsupported claims are withheld.")
            run.italic = True
            run.font.size = Pt(8)

    # Key Result Metric
    top_val = numbers.get("top_value")
    pearson_r = numbers.get("pearson_r")
    if top_val is not None or pearson_r is not None:
        val_str = _fmt_compact(top_val) if top_val is not None else f"r = {float(pearson_r):.3f}"
        p_metric = doc.add_paragraph()
        r_m = p_metric.add_run(f"{val_str}\n")
        r_m.bold = True
        r_m.font.size = Pt(19)
        r_m.font.color.rgb = RGBColor(217, 119, 6)
        p_metric.add_run(f"Category: {kind}  |  Observations: {numbers.get('n', 'N/A')}").italic = True

    # Results Table
    table_meta = _extract_table_data(kind, numbers, str(investigation_data.get("question") or ""))
    if table_meta:
        doc.add_heading(table_meta["title"], level=2)
        headers = table_meta["headers"]
        rows = table_meta["rows"]
        t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.LEFT

        # Header Row
        hdr_cells = t.rows[0].cells
        for i, header_text in enumerate(headers):
            hdr_cells[i].text = header_text
            for p in hdr_cells[i].paragraphs:
                for r in p.runs:
                    r.bold = True
                    r.font.size = Pt(9.5)

        # Data Rows
        for row_idx, row_data in enumerate(rows):
            row_cells = t.rows[row_idx + 1].cells
            for col_idx, cell_value in enumerate(row_data):
                row_cells[col_idx].text = str(cell_value)
                for p in row_cells[col_idx].paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9.5)

    # Interpretation, calculation transparency, and visual evidence
    if report_contract.get("interpreted"):
        doc.add_heading("What MindEd interpreted", level=2)
        for item in report_contract["interpreted"]:
            doc.add_paragraph(str(item), style="List Bullet")

    if report_contract.get("calculation"):
        doc.add_heading("How the result was calculated", level=2)
        for item in report_contract["calculation"]:
            doc.add_paragraph(str(item), style="List Bullet")

    report_charts = report_contract.get("visualizations") or []
    if report_charts:
        doc.add_heading("Visualizations", level=2)
        for chart in report_charts[:4]:
            doc.add_heading(str(chart.get("title") or "Visualization"), level=3)
            if chart.get("explanation"):
                doc.add_paragraph(str(chart["explanation"]))
            png = _render_report_chart_png(chart)
            if png:
                doc.add_picture(BytesIO(png), width=Inches(6.1))

    # Key Insights
    details = analyst_result.get("details") or []
    if details:
        doc.add_heading("Key Insights", level=2)
        for d in details[:3]:
            doc.add_paragraph(d, style="List Bullet")

    # Limitations & Assumptions
    caveats = analyst_result.get("caveats") or []
    narrative_caveats = _as_dict(report_contract.get("narrative")).get("limitations") or []
    caveats = list(dict.fromkeys([*caveats, *narrative_caveats]))
    if caveats or kind in ("ASSOCIATION", "CORRELATION"):
        doc.add_heading("Limitations & Assumptions", level=2)
        all_caveats = list(caveats)
        if kind in ("ASSOCIATION", "CORRELATION") and not any("causation" in c.lower() for c in all_caveats):
            all_caveats.append("The analysis establishes statistical association, not mechanistic causation.")
        for c in all_caveats:
            doc.add_paragraph(c, style="List Bullet")

    # Recommended Actions
    next_steps = analyst_result.get("next_steps") or []
    if next_steps:
        doc.add_heading("Recommended Actions", level=2)
        for act in next_steps:
            doc.add_paragraph(act, style="List Number")

    # Evidence & Provenance
    evidence = investigation_data.get("evidence") or []
    manifest_hash = investigation_data.get("reproducible_manifest_hash") or "N/A"
    doc.add_heading("Evidence Audit & Provenance", level=2)
    p_prov = doc.add_paragraph()
    p_prov.add_run(f"Canonical Digest (SHA-256): {canonical_hash}\nManifest Hash: {manifest_hash}\nDocument Seal: {fingerprint_code}").italic = True

    for ev in evidence[:5]:
        stmt = ev.get("statement", "Evidence Item")
        v_status = ev.get("validation_status", "UNVERIFIED")
        calc = ev.get("calculation_summary", "")
        p_ev = doc.add_paragraph(f"{stmt} [{v_status}]", style="List Bullet")
        if calc:
            p_ev.add_run(f"\nCalculation: {calc}").italic = True

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
