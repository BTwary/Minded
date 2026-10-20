"""Cumulative PDF/DOCX exports for a multi-question analysis session.

Every included question must already have a saved, accepted, verified execution
record. The session report is an assembly of those authoritative records; it never
re-runs calculations or invents cross-question aggregates.
"""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Dict, Iterable, List, Optional
import html


def _as_dict(value: Any) -> Dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _str(value: Any, fallback: str = "Not recorded") -> str:
    if value is None or value == "":
        return fallback
    return str(value)


def _question_answer(item: Dict[str, Any]) -> str:
    return _str(item.get("direct_answer") or item.get("main_finding") or _as_dict(item.get("analyst_result")).get("headline"), "No concise answer was saved.")


def _filters_for(item: Dict[str, Any]) -> List[str]:
    contract = _as_dict(item.get("report_contract"))
    result = _as_dict(item.get("analyst_result"))
    for source in (_as_dict(item.get("execution_record")), contract, result, _as_dict(result.get("numbers")), _as_dict(item.get("analysis_plan"))):
        filters = source.get("filters") or source.get("filter_conditions")
        if isinstance(filters, list):
            out = []
            for row in filters:
                if isinstance(row, dict):
                    out.append(f"{row.get('column', '?')} {row.get('operator', '?')} {row.get('value', '')}".strip())
                else:
                    out.append(str(row))
            if out:
                return out
    return []


def _result_rows(item: Dict[str, Any]) -> List[tuple[str, Any, Any]]:
    analyst = _as_dict(item.get("analyst_result"))
    numbers = _as_dict(analyst.get("numbers"))
    table = numbers.get("table") or numbers.get("grouped_values") or item.get("result_table")
    rows: List[tuple[str, Any, Any]] = []
    if isinstance(table, dict):
        for key, value in table.items():
            if isinstance(value, dict):
                rows.append((str(key), value.get("value", value.get("mean", value.get("count", ""))), value.get("n", "")))
            else:
                rows.append((str(key), value, ""))
    elif isinstance(table, list):
        for idx, value in enumerate(table):
            if isinstance(value, dict):
                label = value.get("group", value.get("label", value.get("region", value.get("category", idx + 1))))
                metric = value.get("value", value.get("metric_value", value.get("y", "")))
                n = value.get("n", value.get("count", ""))
                rows.append((str(label), metric, n))
    return rows


def _visualizations(item: Dict[str, Any]) -> List[Dict[str, Any]]:
    for source in (_as_dict(item.get("report_contract")), _as_dict(item.get("analyst_result"))):
        values = source.get("visualizations") or source.get("charts")
        if isinstance(values, list):
            return [v for v in values if isinstance(v, dict)]
    return []


def _chart_png(chart: Dict[str, Any]) -> Optional[bytes]:
    """Render saved chart data to a static image when its schema is understood."""
    data = chart.get("data")
    if not isinstance(data, list) or not data:
        return None
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        chart_type = str(chart.get("chart_type") or chart.get("type") or "").casefold()
        x_field = _as_dict(chart.get("x_axis")).get("field") or _as_dict(chart.get("x")).get("field")
        y_field = _as_dict(chart.get("y_axis")).get("field") or _as_dict(chart.get("y")).get("field")
        first = data[0] if isinstance(data[0], dict) else {}
        x_field = x_field or next((k for k in first if k not in {"value", "y", "n", "count"}), None)
        y_field = y_field or next((k for k in first if k in {"value", "y", "metric_value", "count", "mean", "sum"}), None)
        if not x_field or not y_field or not all(isinstance(row, dict) for row in data):
            return None
        xs = [row.get(x_field) for row in data]
        ys = [row.get(y_field) for row in data]
        fig, ax = plt.subplots(figsize=(7.2, 3.8), constrained_layout=True)
        if chart_type in {"line", "trend", "timeseries", "time_series"}:
            ax.plot(xs, ys, marker="o")
        elif chart_type in {"scatter", "scatterplot"}:
            ax.scatter(xs, ys, alpha=0.75)
        elif chart_type in {"waterfall"}:
            starts = [row.get("start", 0) for row in data]
            ends = [row.get("end", row.get(y_field, 0)) for row in data]
            ax.bar(range(len(data)), [float(e or 0) - float(s or 0) for s, e in zip(starts, ends)], bottom=starts)
            ax.set_xticks(range(len(data)), [str(x) for x in xs], rotation=25, ha="right")
        else:
            try:
                numeric_y = [float(y) for y in ys]
            except (TypeError, ValueError):
                plt.close(fig)
                return None
            ax.bar([str(x) for x in xs], numeric_y)
            ax.tick_params(axis="x", labelrotation=25)
            ax.set_ylim(bottom=0 if min(numeric_y, default=0) >= 0 else None)
        ax.set_title(_str(chart.get("title"), "Saved visualization"))
        ax.set_xlabel(_str(_as_dict(chart.get("x_axis")).get("label") or x_field, ""))
        ax.set_ylabel(_str(_as_dict(chart.get("y_axis")).get("label") or y_field, ""))
        image = BytesIO()
        fig.savefig(image, format="png", dpi=150)
        plt.close(fig)
        return image.getvalue()
    except Exception:
        return None


def export_analysis_session_to_pdf(session_data: Dict[str, Any], questions: List[Dict[str, Any]]) -> bytes:
    """Create a branded cover and append the already certified question reports."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from pypdf import PdfReader, PdfWriter
    from packages.analytics_core.src.reporting.exporter import export_investigation_to_pdf

    cover_bytes = BytesIO()
    doc = SimpleDocTemplate(cover_bytes, pagesize=letter, rightMargin=.7*inch, leftMargin=.7*inch, topMargin=.65*inch, bottomMargin=.65*inch)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="MindEdTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=23, leading=28, textColor=colors.HexColor("#111827"), alignment=TA_LEFT, spaceAfter=10))
    styles.add(ParagraphStyle(name="MindEdSection", parent=styles["Heading2"], fontSize=12, leading=15, textColor=colors.HexColor("#B45309"), spaceBefore=12, spaceAfter=5))
    styles.add(ParagraphStyle(name="MindEdBody", parent=styles["BodyText"], fontSize=9, leading=13, textColor=colors.HexColor("#334155")))
    story = [
        Paragraph("MINDED AA-OS · CUMULATIVE ANALYSIS REPORT", styles["MindEdSection"]),
        Spacer(1, 8),
        Paragraph(html.escape(_str(session_data.get("title"), "Analysis report")), styles["MindEdTitle"]),
        Paragraph(f"Status: {html.escape(_str(session_data.get('status'), 'OPEN'))} · Questions included: {len(questions)} · Report revision: {session_data.get('report_revision', 0)}", styles["MindEdBody"]),
        Paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}", styles["MindEdBody"]),
        Spacer(1, 16),
        Paragraph("EXECUTIVE OVERVIEW", styles["MindEdSection"]),
        Paragraph("This report combines accepted and verified question-level analyses from one analysis session. Each question section retains its own executed calculation, evidence, limitations, chart(s), and audit seal. Values are not re-aggregated across questions.", styles["MindEdBody"]),
        Spacer(1, 10),
        Paragraph("QUESTIONS INCLUDED", styles["MindEdSection"]),
    ]
    table_data = [["#", "Question", "Accepted finding"]]
    for idx, item in enumerate(questions, start=1):
        question_text = html.escape(_str(item.get("question"), "Analytical question"))
        answer = html.escape(_question_answer(item))
        table_data.append([str(idx), Paragraph(question_text, styles["MindEdBody"]), Paragraph(answer, styles["MindEdBody"])])
    table = Table(table_data, colWidths=[.35*inch, 3.15*inch, 3.15*inch], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 8),
        ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 6), ("RIGHTPADDING", (0,0), (-1,-1), 6),
        ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]))
    story.extend([table, Spacer(1, 18), Paragraph("AUDIT NOTE", styles["MindEdSection"]), Paragraph("The cover and section order describe the analysis session. Numerical claims remain attributable to the certified question-level exports included after this cover. A question is included only when its latest review status is ACCEPTED_AND_VERIFIED.", styles["MindEdBody"])])
    doc.build(story)

    writer = PdfWriter()
    writer.append(PdfReader(cover_bytes))
    for item in questions:
        writer.append(PdfReader(BytesIO(export_investigation_to_pdf(item))))
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def export_analysis_session_to_docx(session_data: Dict[str, Any], questions: List[Dict[str, Any]]) -> bytes:
    """Build an editable cumulative report from stored accepted question records."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.65)
    section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.75)
    section.right_margin = Inches(.75)
    title = doc.add_heading(_str(session_data.get("title"), "Analysis report"), 0)
    title.runs[0].font.color.rgb = RGBColor(180, 83, 9)
    doc.add_paragraph("MINDED AA-OS · CUMULATIVE ANALYSIS REPORT")
    doc.add_paragraph(f"Status: {_str(session_data.get('status'), 'OPEN')} | Questions included: {len(questions)} | Report revision: {session_data.get('report_revision', 0)}")
    doc.add_paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}")
    doc.add_heading("Executive overview", level=1)
    doc.add_paragraph("This report combines accepted and verified question-level analyses from one analysis session. Each question retains its own interpretation, calculation, filters, evidence, limitations and visualizations. No new calculations were performed during export.")
    doc.add_heading("Contents", level=1)
    for idx, item in enumerate(questions, start=1):
        doc.add_paragraph(f"{idx}. {_str(item.get('question'), 'Analytical question')}", style="List Number")
    for idx, item in enumerate(questions, start=1):
        doc.add_page_break()
        doc.add_heading(f"Question {idx}: {_str(item.get('question'), 'Analytical question')}", level=1)
        doc.add_paragraph(f"Status: {_str(item.get('status'))} | Verdict: {_str(item.get('verdict_type') or item.get('verdict'))} | Review: {_str(item.get('review_status'))}")
        doc.add_heading("Answer", level=2)
        doc.add_paragraph(_question_answer(item))

        contract = _as_dict(item.get("report_contract"))
        narrative = _as_dict(contract.get("narrative"))
        if narrative.get("summary"):
            doc.add_heading(_str(narrative.get("title"), "Interpretation"), level=2)
            doc.add_paragraph(str(narrative["summary"]))

        doc.add_heading("Calculation and population", level=2)
        exec_record = _as_dict(item.get("execution_record"))
        plan = _as_dict(item.get("analysis_plan"))
        doc.add_paragraph(f"Method: {_str(exec_record.get('aggregation') or plan.get('requested_aggregation') or plan.get('method'))}")
        doc.add_paragraph(f"Grouping: {_str(exec_record.get('grouping') or plan.get('group_dimension') or plan.get('grouping_columns'), 'Overall (no group)')}")
        sample_count = exec_record.get("eligible_input_records") or exec_record.get("rows_analyzed") or plan.get("population_row_count")
        doc.add_paragraph(f"Eligible observations: {_str(sample_count)}")
        filters = _filters_for(item)
        doc.add_paragraph("Filters: " + ("; ".join(filters) if filters else "No separately stored filter list was available in the report payload."))

        result_rows = _result_rows(item)
        if result_rows:
            doc.add_heading("Saved results", level=2)
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Shading Accent 1"
            table.rows[0].cells[0].text = "Group"
            table.rows[0].cells[1].text = "Value"
            table.rows[0].cells[2].text = "Observations"
            for group, value, n in result_rows:
                cells = table.add_row().cells
                cells[0].text, cells[1].text, cells[2].text = str(group), str(value), str(n)

        chart_images = []
        for chart in _visualizations(item):
            image = _chart_png(chart)
            if image:
                chart_images.append((chart, image))
        if chart_images:
            doc.add_heading("Visualizations", level=2)
            for chart, image in chart_images:
                doc.add_paragraph(_str(chart.get("title"), "Saved visualization"))
                doc.add_picture(BytesIO(image), width=Inches(6.1))
                explanation = chart.get("explanation")
                if explanation:
                    doc.add_paragraph(str(explanation))
        elif _visualizations(item):
            doc.add_heading("Visualizations", level=2)
            doc.add_paragraph("The stored visualization specification is preserved in the question record, but its data shape was not supported by this static DOCX renderer. See the included PDF section for the rendered chart.")

        findings = item.get("findings")
        if isinstance(findings, list) and findings:
            doc.add_heading("Findings and evidence", level=2)
            for finding in findings:
                if isinstance(finding, dict):
                    text = finding.get("statement") or finding.get("claim") or finding.get("description")
                else:
                    text = finding
                if text:
                    doc.add_paragraph(str(text), style="List Bullet")
        caveats = _as_dict(item.get("analyst_result")).get("caveats") or contract.get("limitations") or []
        if caveats:
            doc.add_heading("Limitations and assumptions", level=2)
            for caveat in caveats if isinstance(caveats, list) else [caveats]:
                doc.add_paragraph(str(caveat), style="List Bullet")
        doc.add_heading("Audit and provenance", level=2)
        doc.add_paragraph(f"Investigation ID: {_str(item.get('id'))}")
        doc.add_paragraph(f"Execution contract: {_str(exec_record.get('contract_match') or exec_record.get('execution_contract_status'))}")
        doc.add_paragraph(f"Verification: {_str(item.get('verification_status') or item.get('review_status'))}")
        provenance = _as_dict(item.get("provenance"))
        if provenance:
            for key in ("dataset_version", "dataset_fingerprint", "manifest_hash", "canonical_digest"):
                if provenance.get(key):
                    doc.add_paragraph(f"{key.replace('_', ' ').title()}: {provenance[key]}")

    output = BytesIO()
    doc.save(output)
    return output.getvalue()
